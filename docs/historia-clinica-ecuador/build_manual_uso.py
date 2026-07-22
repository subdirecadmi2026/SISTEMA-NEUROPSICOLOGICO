#!/usr/bin/env python3
"""Genera gráficas densas y el Manual de Uso PROT-HC-001-MU (Word)."""

from __future__ import annotations

import os
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips

ROOT = Path(__file__).resolve().parent
OUT_IMG = ROOT / "manual_img"
OUT_DOCX = ROOT / "MANUAL_USO_ACTAS_ANEXOS_PROT-HC-001.docx"
LOGO = ROOT / "logo_msp.png"
ARTIFACTS = Path("/opt/cursor/artifacts")

NAVY = (28, 58, 92)
TEAL = (46, 125, 132)
GREEN = (46, 125, 80)
BROWN = (140, 90, 55)
GOLD = (180, 130, 40)
ORANGE = (170, 95, 45)
MAROON = (130, 55, 60)
SOFT = (245, 248, 250)
LINE = (210, 220, 228)
TEXT = (32, 42, 52)
MUTED = (90, 105, 118)
WHITE = (255, 255, 255)

FONT_REG = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
FONT_BOLD = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"


def F(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def wrap(draw: ImageDraw.ImageDraw, text: str, font, max_w: int) -> list[str]:
    words = text.split()
    lines: list[str] = []
    cur = ""
    for w in words:
        trial = (cur + " " + w).strip()
        if draw.textlength(trial, font=font) <= max_w:
            cur = trial
        else:
            if cur:
                lines.append(cur)
            cur = w
    if cur:
        lines.append(cur)
    return lines or [""]


def rounded(draw, box, fill, outline=None, width=2, radius=14):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow_down(draw, x, y1, y2, color=NAVY, tip=14):
    draw.line((x, y1, x, y2 - tip), fill=color, width=5)
    draw.polygon([(x, y2), (x - tip, y2 - tip), (x + tip, y2 - tip)], fill=color)


def header_bar(draw, w, h, title, subtitle=None, title_size=42, sub_size=22):
    draw.rectangle((0, 0, w, h), fill=NAVY)
    ft = F(title_size, True)
    bb = draw.textbbox((0, 0), title, font=ft)
    tw = bb[2] - bb[0]
    th = bb[3] - bb[1]
    if subtitle:
        fs = F(sub_size, False)
        sb = draw.textbbox((0, 0), subtitle, font=fs)
        sh = sb[3] - sb[1]
        gap = 10
        total = th + gap + sh
        y0 = (h - total) // 2
        draw.text(((w - tw) // 2, y0), title, font=ft, fill=WHITE)
        sw = sb[2] - sb[0]
        draw.text(((w - sw) // 2, y0 + th + gap), subtitle, font=fs, fill=(200, 220, 230))
    else:
        draw.text(((w - tw) // 2, (h - th) // 2 - 4), title, font=ft, fill=WHITE)


def footer(draw, w, h, y, text):
    draw.text((28, y), text, font=F(18), fill=MUTED)
    return h  # caller sets canvas to content


# ---------------------------------------------------------------------------
# GRAPHICS
# ---------------------------------------------------------------------------

def build_cuando_usar(path: Path):
    # Narrower canvas → larger text when scaled to A4 width (~19 cm)
    W = 1600
    pad = 22
    header_h = 100
    blocks = [
        ("1. SE DETECTA PÉRDIDA / EXTRAVÍO / DETERIORO DE HCU",
         "Inicie el protocolo de inmediato", TEAL, False),
        ("ACTA 01 — Detección y notificación",
         "Usar YA (≤ 30 min). Contener el área y notificar a Admisiones / Calidad", NAVY, True),
        ("2. ACTIVAR PLAN + COMITÉS (Art. 13)",
         "Código INC-HC-AAAA-### · Clasificar E1–E5 · ≤ 2 horas", (55, 90, 120), False),
        ("ACTA 02 — Diagnóstico e inventario",
         "Autorizar ingreso al archivo · Conformar archivo recuperado", (70, 110, 100), True),
    ]
    box_h = 108
    gap_arr = 34
    decision_h = 100
    branch_h = 170
    parallel_h = 120
    end_h = 108
    H = (
        header_h + 18
        + len(blocks) * (box_h + gap_arr)
        + decision_h + gap_arr
        + branch_h + gap_arr
        + parallel_h + gap_arr
        + 2 * (end_h + gap_arr)
        + 60
    )
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    header_bar(d, W, header_h,
               "¿CUÁNDO USAR CADA ACTA O ANEXO?",
               "Hospital General Puyo · PROT-HC-001 · Siga las flechas",
               34, 20)

    x0, x1 = pad, W - pad
    cx = W // 2
    y = header_h + 18
    ft_t = F(28, True)
    ft_s = F(22)

    def step_box(y, title, sub, color, accent=False):
        box = (x0, y, x1, y + box_h)
        rounded(d, box, fill=color if accent else (248, 251, 252),
                outline=color, width=4, radius=14)
        if accent:
            d.text((x0 + 20, y + 18), title, font=ft_t, fill=WHITE)
            d.text((x0 + 20, y + 62), sub, font=ft_s, fill=(230, 240, 245))
        else:
            d.rectangle((x0, y, x0 + 14, y + box_h), fill=color)
            d.text((x0 + 28, y + 18), title, font=ft_t, fill=NAVY)
            d.text((x0 + 28, y + 62), sub, font=ft_s, fill=MUTED)
        return y + box_h

    for title, sub, color, accent in blocks:
        y = step_box(y, title, sub, color, accent)
        arrow_down(d, cx, y + 2, y + gap_arr - 2, tip=12)
        y += gap_arr

    dw, dh = 560, decision_h
    dx0 = cx - dw // 2
    rounded(d, (dx0, y, dx0 + dw, y + dh), fill=WHITE, outline=NAVY, width=5, radius=16)
    q = "¿SE RECUPERA LA HCU?"
    qb = d.textbbox((0, 0), q, font=F(30, True))
    d.text((cx - (qb[2] - qb[0]) // 2, y + (dh - (qb[3] - qb[1])) // 2 - 2),
           q, font=F(30, True), fill=NAVY)
    y_dec_bottom = y + dh

    # SÍ / NO badges beside decision
    rounded(d, (pad, y + 22, pad + 90, y + 78), fill=GREEN, radius=10)
    d.text((pad + 22, y + 34), "SÍ", font=F(28, True), fill=WHITE)
    rounded(d, (W - pad - 90, y + 22, W - pad, y + 78), fill=BROWN, radius=10)
    d.text((W - pad - 68, y + 34), "NO", font=F(28, True), fill=WHITE)

    y_branch = y_dec_bottom + gap_arr
    left_box = (pad, y_branch, cx - 12, y_branch + branch_h)
    right_box = (cx + 12, y_branch, W - pad, y_branch + branch_h)
    lx = (left_box[0] + left_box[2]) // 2
    rx = (right_box[0] + right_box[2]) // 2
    d.line((cx, y_dec_bottom, cx, y_dec_bottom + 12), fill=NAVY, width=4)
    d.line((lx, y_dec_bottom + 12, rx, y_dec_bottom + 12), fill=NAVY, width=4)
    d.line((lx, y_dec_bottom + 12, lx, y_branch), fill=GREEN, width=5)
    d.line((rx, y_dec_bottom + 12, rx, y_branch), fill=BROWN, width=5)

    rounded(d, left_box, fill=GREEN, outline=GREEN, width=3, radius=14)
    d.text((left_box[0] + 14, left_box[1] + 16),
           "ACTA 03 — Recuperación", font=F(24, True), fill=WHITE)
    for i, line in enumerate(wrap(d,
            "Firmar: Servicio + Admisiones + Calidad. Arts. 18–19 si hay pérdida parcial/total.",
            F(20), int(left_box[2] - left_box[0] - 28))):
        d.text((left_box[0] + 14, left_box[1] + 58 + i * 28), line, font=F(20), fill=(230, 245, 235))

    rounded(d, right_box, fill=BROWN, outline=BROWN, width=3, radius=14)
    d.text((right_box[0] + 14, right_box[1] + 16),
           "ACTA 04 — Reconstrucción", font=F(24, True), fill=WHITE)
    for i, line in enumerate(wrap(d,
            "Solo datos verificables · Declarar lagunas. Prohibido inventar información clínica.",
            F(20), int(right_box[2] - right_box[0] - 28))):
        d.text((right_box[0] + 14, right_box[1] + 58 + i * 28), line, font=F(20), fill=(245, 235, 225))

    y = y_branch + branch_h
    mid = y + gap_arr // 2
    d.line((lx, y, lx, mid), fill=TEAL, width=4)
    d.line((rx, y, rx, mid), fill=TEAL, width=4)
    d.line((lx, mid, rx, mid), fill=TEAL, width=4)
    arrow_down(d, cx, mid, y + gap_arr - 2, TEAL, tip=12)
    y += gap_arr

    box = (x0, y, x1, y + parallel_h)
    rounded(d, box, fill=(232, 244, 248), outline=TEAL, width=5, radius=14)
    d.text((x0 + 18, y + 18), "EN PARALELO → ANEXO A (si hay paciente en atención)",
           font=F(24, True), fill=TEAL)
    for i, line in enumerate(wrap(d,
            "Formulario de contingencia para NO interrumpir la atención mientras se recupera o reconstruye.",
            F(20), x1 - x0 - 36)):
        d.text((x0 + 18, y + 60 + i * 28), line, font=F(20), fill=MUTED)
    y += parallel_h
    arrow_down(d, cx, y + 2, y + gap_arr - 2, tip=12)
    y += gap_arr

    y = step_box(y, "ACTA 05 — Notificación al usuario (si aplica)",
                 "Entorno privado · Copia certificada ≤ 48 h si la solicita", (55, 95, 130), True)
    arrow_down(d, cx, y + 2, y + gap_arr - 2, tip=12)
    y += gap_arr
    y = step_box(y, "ACTA 06 — Informe de cierre",
                 "Medidas correctivas · Archivo INC-HC restringido · ≤ 5 días hábiles",
                 NAVY, True)
    y += 14
    d.text((16, y), "PROT-HC-001 · Manual de uso · Hospital General Puyo · 0991976454",
           font=F(16), fill=MUTED)
    y += 36
    img = img.crop((0, 0, W, min(y + 4, H)))
    img.save(path, "PNG", optimize=True)
    print("wrote", path, img.size)


def build_quien_firma(path: Path):
    W = 1600
    header_h = 96
    rows = [
        ("Acta 01", "✓ Firma", "✓ Jefatura", "—", "—", "✓ Detecta"),
        ("Acta 02", "✓ Firma", "✓ Firma", "✓ Firma", "✓ Firma", "TI si aplica"),
        ("Acta 03", "✓ Firma", "✓ Firma", "✓ Firma", "—", "—"),
        ("Acta 04", "✓ Firma", "✓ Firma", "✓ Firma", "✓ Aprueba", "✓ Prof."),
        ("Acta 05", "✓ Firma", "—", "✓ Vo.Bo.", "—", "✓ Usuario"),
        ("Acta 06", "✓ Elabora", "✓ Adic.", "✓ Firma", "✓ Aprueba", "TH si aplica"),
        ("Anexo A", "✓ Recibe", "—", "✓ Firma", "—", "✓ Prof."),
        ("Anexo B", "✓ Verifica", "—", "✓ Verifica", "✓ Verifica", "—"),
    ]
    cols = ["Documento", "Admisiones", "Servicio", "Calidad", "Dirección", "Otros"]
    legend = [
        (TEAL, "Admisiones"),
        (GREEN, "Servicio"),
        ((70, 100, 125), "Calidad"),
        (NAVY, "Dirección"),
        (BROWN, "Otros"),
    ]
    row_h = 72
    hdr_h = 52
    pad = 16
    legend_h = 56
    highlight_h = 110
    table_top = header_h + 12 + legend_h + 10
    H = table_top + hdr_h + len(rows) * row_h + 14 + highlight_h + 50
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    header_bar(d, W, header_h,
               "¿QUIÉN FIRMA CADA DOCUMENTO?",
               "Matriz de responsables — Hospital General Puyo",
               32, 18)

    y = header_h + 12
    x = pad
    for color, label in legend:
        tw = int(d.textlength(label, font=F(18, True))) + 28
        rounded(d, (x, y, x + tw, y + 44), fill=color, radius=10)
        d.text((x + 14, y + 10), label, font=F(18, True), fill=WHITE)
        x += tw + 10

    col_w = [250, 250, 250, 250, 250, 250]
    usable = W - 2 * pad
    scale = usable / sum(col_w)
    col_w = [int(c * scale) for c in col_w]
    # fix rounding
    col_w[-1] += usable - sum(col_w)
    y = table_top
    x = pad
    for i, c in enumerate(cols):
        box = (x, y, x + col_w[i], y + hdr_h)
        d.rectangle(box, fill=NAVY)
        font = F(18, True)
        tw = d.textlength(c, font=font)
        d.text((x + (col_w[i] - tw) / 2, y + 14), c, font=font, fill=WHITE)
        x += col_w[i]
    y += hdr_h

    for ri, row in enumerate(rows):
        x = pad
        bg = (248, 250, 252) if ri % 2 else WHITE
        if row[0] == "Acta 03":
            bg = (232, 245, 236)
        for i, cell in enumerate(row):
            box = (x, y, x + col_w[i], y + row_h)
            d.rectangle(box, fill=bg, outline=LINE, width=2)
            color = NAVY if i == 0 else (GREEN if "✓" in cell else MUTED)
            font = F(22, True) if i == 0 or "✓" in cell else F(20)
            tw = d.textlength(cell, font=font)
            d.text((x + (col_w[i] - tw) / 2, y + 22), cell, font=font, fill=color)
            x += col_w[i]
        y += row_h

    y += 14
    rounded(d, (pad, y, W - pad, y + highlight_h), fill=(230, 245, 234),
            outline=GREEN, width=4, radius=12)
    lines = [
        "DESTACADO — ACTA 03",
        "Firmas obligatorias = Servicio + Admisiones + Calidad",
    ]
    for i, line in enumerate(lines):
        font = F(26, True)
        tw = d.textlength(line, font=font)
        d.text(((W - tw) / 2, y + 22 + i * 36), line, font=font, fill=GREEN)
    y += highlight_h + 12
    d.text((pad, y), "✓ = firma obligatoria   ·   — = no aplica   ·   Vo.Bo. = visto bueno",
           font=F(16), fill=MUTED)
    y += 36
    img = img.crop((0, 0, W, y))
    img.save(path, "PNG", optimize=True)
    print("wrote", path, img.size)


def build_como_llenar(path: Path):
    """Filas horizontales a ancho completo → tipografía grande en A4."""
    W = 1600
    header_h = 100
    rules = [
        ("1", "IDENTIFIQUE",
         "Escriba el N.º INC-HC-AAAA-###, fecha, hora y datos del usuario o de la HCU. Sin omitir campos obligatorios."),
        ("2", "MARQUE",
         "Use casillas ☑ solo en la opción correcta. No deje opciones ambiguas ni marcas dobles."),
        ("3", "DESCRIBA",
         "Redacte con letra clara, tinta azul o negra. Solo hechos verificables. Sin tachones ni corrector líquido."),
        ("4", "ADJUNTE",
         "Anexe evidencias: inventarios, copias, Anexo A, resultados, respaldos. Numere los anexos."),
        ("5", "FIRME",
         "Cada responsable firma con nombre, cargo y fecha. Está prohibido dejar firmas en blanco."),
        ("6", "ARCHIVE",
         "Guarde el original en el expediente INC-HC (área restringida). Marque CONFIDENCIAL."),
    ]
    extras = [
        "Tinta azul o negra; no use lápiz ni corrector líquido.",
        "Error: tache con una línea, escriba “error” y firme al margen.",
        "Campos no aplicables: escriba “N/A”.",
        "Original al expediente INC-HC (CONFIDENCIAL); conserve copia legible.",
    ]
    pad = 18
    row_h = 118
    gap = 10
    extras_h = 48 + len(extras) * 36
    H = header_h + 14 + len(rules) * (row_h + gap) + 8 + extras_h + 40
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    header_bar(d, W, header_h,
               "CÓMO LLENAR CORRECTAMENTE LAS ACTAS",
               "Seis reglas de oro — escritura legible, completa y válida",
               30, 18)

    y = header_h + 14
    for num, title, body in rules:
        rounded(d, (pad, y, W - pad, y + row_h), fill=SOFT, outline=TEAL, width=3, radius=12)
        d.ellipse((pad + 16, y + 28, pad + 72, y + 84), fill=TEAL)
        nb = d.textbbox((0, 0), num, font=F(28, True))
        d.text((pad + 44 - (nb[2] - nb[0]) // 2, y + 40), num, font=F(28, True), fill=WHITE)
        d.text((pad + 90, y + 18), title, font=F(28, True), fill=NAVY)
        for li, line in enumerate(wrap(d, body, F(22), W - pad * 2 - 110)):
            d.text((pad + 90, y + 58 + li * 28), line, font=F(22), fill=TEXT)
        y += row_h + gap

    y += 4
    rounded(d, (pad, y, W - pad, y + extras_h), fill=(255, 250, 240), outline=GOLD, width=3, radius=12)
    d.text((pad + 18, y + 10), "REGLAS ADICIONALES", font=F(22, True), fill=GOLD)
    for i, e in enumerate(extras):
        d.text((pad + 18, y + 48 + i * 36), f"•  {e}", font=F(20), fill=TEXT)
    y += extras_h + 12
    d.text((pad, y), "PROT-HC-001 · Manual de uso · Hospital General Puyo", font=F(16), fill=MUTED)
    y += 34
    img = img.crop((0, 0, W, y))
    img.save(path, "PNG", optimize=True)
    print("wrote", path, img.size)


def build_senaletica(path: Path):
    """Filas E1–E5 apiladas: máxima legibilidad en página A4."""
    W = 1600
    header_h = 100
    events = [
        (TEAL, "E1", "Extravío unitario", "01 → 02 → 03 (o 04 si no aparece)  →  luego 05 → 06"),
        (GREEN, "E2", "Deterioro parcial", "01 → 02 → 03/04 según el daño  →  luego 05 → 06"),
        (GOLD, "E3", "Pérdida parcial", "01 → 02 → 03 + archivo recuperado  →  luego 05 → 06"),
        (ORANGE, "E4", "Pérdida total", "01 → 02 → 03/04 (Arts. 18–19)  →  luego 05 → 06"),
        (MAROON, "E5", "Falla electrónica", "01 → 02 → 03 + Tecnología/backup  →  luego 05 → 06"),
    ]
    pad = 16
    row_h = 100
    gap = 10
    remember_h = 130
    H = header_h + 14 + len(events) * (row_h + gap) + 10 + remember_h + 40
    img = Image.new("RGB", (W, H), WHITE)
    d = ImageDraw.Draw(img)
    header_bar(d, W, header_h,
               "SEÑALÉTICA RÁPIDA POR TIPO DE EVENTO (E1–E5)",
               "Qué actas priorizar según la clasificación del incidente",
               28, 18)

    y = header_h + 14
    for color, code, title, seq in events:
        rounded(d, (pad, y, W - pad, y + row_h), fill=WHITE, outline=color, width=4, radius=12)
        d.rectangle((pad, y, pad + 110, y + row_h), fill=color)
        cb = d.textbbox((0, 0), code, font=F(32, True))
        d.text((pad + 55 - (cb[2] - cb[0]) // 2, y + 30), code, font=F(32, True), fill=WHITE)
        d.text((pad + 128, y + 16), title, font=F(26, True), fill=NAVY)
        d.text((pad + 128, y + 56), seq, font=F(22, True), fill=color)
        y += row_h + gap

    y += 4
    rounded(d, (pad, y, W - pad, y + remember_h), fill=(232, 244, 248), outline=TEAL, width=4, radius=12)
    d.text((pad + 18, y + 14), "RECUERDE", font=F(24, True), fill=TEAL)
    for i, line in enumerate([
        "Anexo A: siempre que haya paciente en atención y la HCU no esté disponible.",
        "Anexo B: verifica la implementación del protocolo en el hospital (NO es por cada incidente).",
    ]):
        d.text((pad + 18, y + 52 + i * 34), line, font=F(20), fill=TEXT)
    y += remember_h + 10
    d.text((pad, y), "PROT-HC-001 · Hospital General Puyo · MSP Ecuador · 0991976454",
           font=F(16), fill=MUTED)
    y += 34
    img = img.crop((0, 0, W, y))
    img.save(path, "PNG", optimize=True)
    print("wrote", path, img.size)


# ---------------------------------------------------------------------------
# WORD MANUAL
# ---------------------------------------------------------------------------

def set_run(run, size=11, bold=False, color=NAVY, font="Calibri"):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def shade_cell(cell, hex_color: str):
    tc = cell._te if hasattr(cell, "_te") else cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        if edge in kwargs:
            tag = OxmlElement(f"w:{edge}")
            tag.set(qn("w:val"), kwargs[edge].get("val", "single"))
            tag.set(qn("w:sz"), str(kwargs[edge].get("sz", 8)))
            tag.set(qn("w:color"), kwargs[edge].get("color", "1C3A5C"))
            tcBorders.append(tag)
    tcPr.append(tcBorders)


def tight_para(p, before=0, after=2, line=1.0):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line


def add_heading_bar(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "1C3A5C")
    p = cell.paragraphs[0]
    tight_para(p, 0, 0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run(run, 12, True, WHITE, "Calibri")
    # reduce cell margins
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for m, v in (("top", "40"), ("bottom", "40"), ("left", "80"), ("right", "80")):
        el = OxmlElement(f"w:{m}")
        el.set(qn("w:w"), v)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)
    # spacer tiny
    sp = doc.add_paragraph()
    tight_para(sp, 0, 2)


def add_callout(doc, text, fill="E8F4F8", border="2E7D84"):
    table = doc.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    shade_cell(cell, fill)
    set_cell_border(cell,
                    top={"sz": 12, "color": border},
                    bottom={"sz": 12, "color": border},
                    left={"sz": 24, "color": border},
                    right={"sz": 12, "color": border})
    p = cell.paragraphs[0]
    tight_para(p, 0, 0)
    run = p.add_run(text)
    set_run(run, 10, True, NAVY)
    sp = doc.add_paragraph()
    tight_para(sp, 0, 2)


def add_image_full(doc, path: Path, width_cm=19.2):
    p = doc.add_paragraph()
    tight_para(p, 0, 4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm))


def add_caption(doc, text):
    p = doc.add_paragraph()
    tight_para(p, 0, 4)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run(run, 9, False, MUTED)


def build_docx(img_paths: dict[str, Path]):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(0.7)
    sec.bottom_margin = Cm(0.7)
    sec.left_margin = Cm(0.9)
    sec.right_margin = Cm(0.9)

    # --- Portada compacta ---
    if LOGO.exists():
        t = doc.add_table(rows=1, cols=3)
        t.autofit = True
        c0, c1, c2 = t.rows[0].cells
        p = c0.paragraphs[0]
        tight_para(p)
        r = p.add_run()
        r.add_picture(str(LOGO), width=Cm(2.2))
        p1 = c1.paragraphs[0]
        tight_para(p1)
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p1.add_run("MINISTERIO DE SALUD PÚBLICA DEL ECUADOR")
        set_run(r, 10, True, NAVY)
        p2 = c1.add_paragraph()
        tight_para(p2)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run("HOSPITAL GENERAL PUYO")
        set_run(r, 14, True, TEAL)
        # empty right for balance
        c2.paragraphs[0].text = ""

    title = doc.add_paragraph()
    tight_para(title, 4, 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("MANUAL DE USO DE ACTAS Y ANEXOS")
    set_run(r, 16, True, NAVY, "Georgia")

    sub = doc.add_paragraph()
    tight_para(sub, 0, 2)
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Protocolo de Pérdida de Historia Clínica — PROT-HC-001")
    set_run(r, 11, False, MUTED)

    meta = doc.add_paragraph()
    tight_para(meta, 0, 4)
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Código: PROT-HC-001-MU  ·  Versión 1.1  ·  22/07/2026  ·  CONFIDENCIAL")
    set_run(r, 9, True, NAVY)

    # Contacto strip
    t = doc.add_table(rows=1, cols=3)
    for i, (k, v) in enumerate([
        ("CONTACTO", "0991976454"),
        ("ÁMBITO", "Hospital General Puyo"),
        ("BASE", "PROT-HC-001"),
    ]):
        cell = t.cell(0, i)
        shade_cell(cell, "F2F7F8")
        p = cell.paragraphs[0]
        tight_para(p)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{k}: ")
        set_run(r, 9, True, TEAL)
        r = p.add_run(v)
        set_run(r, 9, False, TEXT)

    add_heading_bar(doc, "CONTROL DE APROBACIÓN")
    t = doc.add_table(rows=1, cols=3)
    for i, (role, name) in enumerate([
        ("ELABORADO POR", "Tlgo. Andrés Romero"),
        ("REVISADO POR", "Dra. Stephany Tayupanda"),
        ("APROBADO POR", "Dr. Gabriel García"),
    ]):
        cell = t.cell(0, i)
        shade_cell(cell, "F7FAFC")
        p = cell.paragraphs[0]
        tight_para(p)
        r = p.add_run(role + "\n")
        set_run(r, 8, True, TEAL)
        r = p.add_run(name + "\n")
        set_run(r, 10, True, NAVY)
        r = p.add_run("Firma: __________  Fecha: ____")
        set_run(r, 8, False, MUTED)

    add_heading_bar(doc, "1. OBJETIVO")
    p = doc.add_paragraph()
    tight_para(p, 0, 4)
    r = p.add_run(
        "Este Manual indica cuándo usar, cómo completar y quién firma las Actas 01–06 y los "
        "Anexos A–B del Protocolo PROT-HC-001 en el Hospital General Puyo. Incluye gráficas "
        "de flujo, matriz de firmas y señalética E1–E5."
    )
    set_run(r, 10, False, TEXT)

    add_heading_bar(doc, "2. ¿CUÁNDO USAR CADA ACTA O ANEXO?")
    add_caption(doc, "Siga las flechas en orden. Cada casilla indica el documento y la acción.")
    add_image_full(doc, img_paths["cuando"], 19.2)
    add_callout(doc, "SEÑAL CLAVE  ·  ACTA 01 siempre primero → ACTA 02 inventario → decisión 03/04 → (Anexo A en paralelo) → 05 → 06")

    add_heading_bar(doc, "3. ¿QUIÉN FIRMA CADA DOCUMENTO?")
    add_caption(doc, "Matriz de firmas obligatorias. ✓ = firma requerida.")
    add_image_full(doc, img_paths["firma"], 19.2)
    add_callout(doc,
                "ACTA 03  ·  Firmas obligatorias: Responsable del servicio + Admisiones + Calidad y Seguridad del Paciente",
                fill="E6F6EA", border="2E7D50")

    add_heading_bar(doc, "4. CÓMO LLENAR CORRECTAMENTE")
    add_image_full(doc, img_paths["llenar"], 19.2)

    add_heading_bar(doc, "5. SEÑALÉTICA RÁPIDA POR EVENTO (E1–E5)")
    add_image_full(doc, img_paths["senal"], 19.2)

    add_heading_bar(doc, "6. FICHAS DE USO POR DOCUMENTO")

    fichas = [
        ("ACTA 01 — Detección y notificación",
         "Al detectar pérdida/extravío/deterioro.",
         "Quien detecta · Jefatura del servicio · Admisiones",
         "≤ 30 minutos",
         ["Anote fecha, hora y solicite/use INC-HC.",
          "Marque E1–E5 y describa el hecho.",
          "Identifique HCU y último movimiento.",
          "Registre contención y notificaciones.",
          "Obtenga las 3 firmas."]),
        ("ACTA 02 — Diagnóstico e inventario",
         "Tras activar el plan y autorizar ingreso al archivo.",
         "Admisiones · Servicio · Calidad · Dirección",
         "1 día hábil (o según magnitud)",
         ["Conforme equipo y registre autorización.",
          "Defina alcance (activo/pasivo/digital).",
          "Inventario B/D/F/C y totales.",
          "Conforme archivo recuperado y confirme E__.",
          "Conclusiones y firmas."]),
        ("ACTA 03 — Recuperación y restauración",
         "Cuando la HCU se localiza, restaura o reintegra.",
         "Responsable del servicio · Admisiones · Calidad",
         "Inmediato tras recuperar",
         ["Indique método (restitución/backup/anexión).",
          "Aplique Arts. 18–19 de numeración.",
          "Resultado por expediente (OK/Parcial/…).",
          "Verifique custodia y nota administrativa.",
          "Las 3 firmas son obligatorias."]),
        ("ACTA 04 — Reconstrucción documental",
         "Si no hay recuperación íntegra o hay lagunas.",
         "Profesional · Admisiones · Calidad · Servicio · Dirección",
         "Tras confirmar no recuperación",
         ["Declare motivo y fuentes verificables.",
          "Prohibido inventar datos; declare lagunas.",
          "Note administrativa con INC-HC.",
          "Adjunte evidencias (Anexo A, resultados…).",
          "Firmas de elaboración y aprobación."]),
        ("ACTA 05 — Notificación al usuario",
         "Si se afectó disponibilidad o acceso del usuario.",
         "Usuario · Funcionario · Admisiones · Vo.Bo. Calidad",
         "Al estabilizar o al solicitarlo",
         ["Informe en entorno privado.",
          "Marque lo comunicado.",
          "Copia certificada ≤ 48 h si solicita.",
          "Firmas de usuario y personal."]),
        ("ACTA 06 — Informe de cierre",
         "Al concluir acciones técnicas y administrativas.",
         "Admisiones (elabora) · Calidad · Dirección",
         "≤ 5 días hábiles",
         ["Resuma causa, alcance y resultados.",
          "Acciones y medidas correctivas con plazos.",
          "Impacto y responsabilidades.",
          "Archive INC-HC en área restringida."]),
        ("ANEXO A — Contingencia asistencial",
         "Paciente en atención sin HCU disponible.",
         "Profesional · Admisiones · Calidad",
         "Inmediato",
         ["Identifique usuario y profesional.",
          "Registre acto clínico completo.",
          "No deje a la vista (CONFIDENCIAL).",
          "Transcriba a la HCU y archive en INC-HC."]),
        ("ANEXO B — Lista de verificación",
         "Verificar implementación del protocolo (no por incidente).",
         "Admisiones · Calidad · Dirección",
         "Trimestral / post-auditoría",
         ["Marque ítems verificados.",
          "Fecha y observaciones.",
          "Tres firmas de verificación."]),
    ]

    for title, cuando, quien, plazo, pasos in fichas:
        # title strip
        t = doc.add_table(rows=1, cols=1)
        cell = t.cell(0, 0)
        shade_cell(cell, "2E7D84")
        p = cell.paragraphs[0]
        tight_para(p)
        r = p.add_run(title)
        set_run(r, 11, True, WHITE)

        meta = doc.add_table(rows=1, cols=3)
        for i, (lab, val) in enumerate([
            ("¿Cuándo?", cuando),
            ("¿Quién firma?", quien),
            ("¿Plazo?", plazo),
        ]):
            cell = meta.cell(0, i)
            shade_cell(cell, "F2F7F8")
            p = cell.paragraphs[0]
            tight_para(p)
            r = p.add_run(lab + "\n")
            set_run(r, 8, True, TEAL)
            r = p.add_run(val)
            set_run(r, 9, False, TEXT)

        p = doc.add_paragraph()
        tight_para(p, 2, 0)
        r = p.add_run("Pasos:")
        set_run(r, 9, True, NAVY)
        for i, step in enumerate(pasos, 1):
            p = doc.add_paragraph()
            tight_para(p, 0, 0)
            r = p.add_run(f"{i}.  {step}")
            set_run(r, 9, False, TEXT)
        sp = doc.add_paragraph()
        tight_para(sp, 0, 3)

    add_heading_bar(doc, "7. ERRORES FRECUENTES")
    err = doc.add_table(rows=1, cols=2)
    hdr = err.rows[0].cells
    for i, h in enumerate(["Error", "Corrección"]):
        shade_cell(hdr[i], "1C3A5C")
        p = hdr[i].paragraphs[0]
        tight_para(p)
        r = p.add_run(h)
        set_run(r, 10, True, WHITE)
    errors = [
        ("Empezar sin Acta 01", "Siempre inicie con detección y notificación."),
        ("Sin código INC-HC", "Asigne INC-HC-AAAA-### al activar el plan."),
        ("Acta 03 sin las 3 firmas", "Servicio + Admisiones + Calidad son obligatorias."),
        ("Reconstruir con datos inventados", "Solo fuentes verificables; declare lagunas (Acta 04)."),
        ("Atender sin Anexo A", "Si no hay HCU y hay paciente, use Anexo A ya."),
        ("Formularios a la vista", "Son CONFIDENCIALES; custodia restringida."),
        ("Cerrar sin Acta 06", "Todo incidente debe cerrarse con informe y medidas."),
    ]
    for a, b in errors:
        row = err.add_row().cells
        for i, txt in enumerate([a, b]):
            p = row[i].paragraphs[0]
            tight_para(p)
            r = p.add_run(txt)
            set_run(r, 9, i == 0, TEXT if i else NAVY)

    add_heading_bar(doc, "8. CHECKLIST RÁPIDO DEL OPERADOR")
    chk = doc.add_table(rows=5, cols=2)
    checks = [
        ("☐  ¿Acta 01 completa y firmada?", "☐  ¿Acta 03 (recuperó) o Acta 04 (no recuperó)?"),
        ("☐  ¿Plan activado y código INC-HC asignado?", "☐  ¿Acta 03 con firmas Servicio + Admisiones + Calidad?"),
        ("☐  ¿Evento clasificado E1–E5?", "☐  ¿Acta 05 al usuario cuando aplica?"),
        ("☐  ¿Acta 02 con inventario / archivo recuperado?", "☐  ¿Acta 06 de cierre y expediente archivado?"),
        ("☐  ¿Anexo A usado si había paciente en atención?", ""),
    ]
    for i, (a, b) in enumerate(checks):
        for j, txt in enumerate([a, b]):
            cell = chk.cell(i, j)
            shade_cell(cell, "F7FAFC")
            p = cell.paragraphs[0]
            tight_para(p)
            r = p.add_run(txt)
            set_run(r, 9, False, TEXT)

    add_callout(doc, "CONTACTO  ·  Dudas de llenado o custodia: 0991976454 · Gestión de Calidad y Seguridad del Paciente / Admisiones")

    add_heading_bar(doc, "APROBACIÓN DEL MANUAL")
    t = doc.add_table(rows=1, cols=3)
    for i, (role, name) in enumerate([
        ("ELABORADO POR", "Tlgo. Andrés Romero"),
        ("REVISADO POR", "Dra. Stephany Tayupanda"),
        ("APROBADO POR", "Dr. Gabriel García"),
    ]):
        cell = t.cell(0, i)
        p = cell.paragraphs[0]
        tight_para(p)
        r = p.add_run(f"{role}\n{name}\n\n__________\nFirma · Fecha")
        set_run(r, 9, False, TEXT)

    OUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_DOCX)
    print("wrote", OUT_DOCX, OUT_DOCX.stat().st_size)


def main():
    OUT_IMG.mkdir(parents=True, exist_ok=True)
    paths = {
        "cuando": OUT_IMG / "grafica_cuando_usar.png",
        "firma": OUT_IMG / "grafica_quien_firma.png",
        "llenar": OUT_IMG / "grafica_como_llenar.png",
        "senal": OUT_IMG / "grafica_senaletica_E.png",
    }
    build_cuando_usar(paths["cuando"])
    build_quien_firma(paths["firma"])
    build_como_llenar(paths["llenar"])
    build_senaletica(paths["senal"])
    build_docx(paths)

    # sync artifacts
    if ARTIFACTS.exists():
        art_img = ARTIFACTS / "manual_img"
        art_img.mkdir(parents=True, exist_ok=True)
        import shutil
        for p in paths.values():
            shutil.copy2(p, art_img / p.name)
        shutil.copy2(OUT_DOCX, ARTIFACTS / OUT_DOCX.name)
        print("synced artifacts")


if __name__ == "__main__":
    main()
